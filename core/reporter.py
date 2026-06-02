from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

DARK_BLUE = "1F3864"
TEAL = "17375E"
LIGHT_GRAY = "F2F2F2"


def _header(ws, row, cells, bg=DARK_BLUE):
    for ci, val in enumerate(cells, 1):
        c = ws.cell(row, ci, val)
        c.fill = PatternFill("solid", fgColor=bg)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _safe_sheet_name(name: str) -> str:
    bad = '/\\?*[]:'
    out = ''.join('_' if c in bad else c for c in name)
    return out[:31]


def export_excel(run_data: dict, out_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    _write_summary(ws, run_data)

    ws_cmp = wb.create_sheet("Comparison")
    _write_comparison(ws_cmp, run_data)

    used = {"Summary", "Comparison"}
    for model_name, model_data in run_data["results"].items():
        if "error" in model_data:
            continue
        sheet = _safe_sheet_name(model_name)
        i = 1
        base = sheet
        while sheet in used:
            sheet = f"{base[:28]}_{i}"
            i += 1
        used.add(sheet)
        ws_m = wb.create_sheet(sheet)
        _write_model_sheet(ws_m, model_name, model_data)

    wb.save(out_path)
    return out_path


def _write_summary(ws, run_data):
    ws.merge_cells("A1:H1")
    c = ws.cell(1, 1, f"Run Summary — Target: {run_data['config']['target']}")
    c.font = Font(bold=True, color="FFFFFF", size=14)
    c.fill = PatternFill("solid", fgColor=DARK_BLUE)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    cfg = run_data["config"]
    panel = run_data["panel"]
    rows = [
        ("Target", cfg["target"]),
        ("Entity / Time", f"{cfg['entity_col']} × {cfg['time_col']}"),
        ("N (full)", panel["n_full"]),
        ("N (after lag)", panel["n_lag"]),
        ("N firms", panel["n_firms"]),
        ("Base lag", cfg["base_lag"]),
        ("Revenue lag", cfg["rev_lag"]),
        ("Winsorize", str(cfg["winsorize_limits"])),
        ("FE",
         f"entity={cfg['entity_effects']} time={cfg['time_effects']} "
         f"industry={cfg['industry_effects']}"),
        ("SE", "Clustered by firm" if cfg["cluster_entity"] else "Robust"),
    ]
    for i, (k, v) in enumerate(rows, 3):
        ws.cell(i, 1, k).font = Font(bold=True)
        ws.cell(i, 2, v)

    next_row = 3 + len(rows) + 1

    # DID setup block
    did = run_data.get("did")
    did_warning = run_data.get("did_warning")
    if did:
        ws.cell(next_row, 1, "DID setup").font = Font(bold=True, color="FFFFFF")
        ws.cell(next_row, 1).fill = PatternFill("solid", fgColor=TEAL)
        ws.merge_cells(start_row=next_row, start_column=1,
                       end_row=next_row, end_column=4)
        next_row += 1
        did_rows = [
            ("Treatment column", did.get("treat_col")),
            ("Cutoff", round(did.get("cutoff", 0), 4)),
            ("Post year (≥)", did.get("post_year")),
            ("Assignment", did.get("assignment", "")),
            ("Firms High", did.get("n_firms_high")),
            ("Firms Low", did.get("n_firms_low")),
            ("Firms dropped (no pre-data)", did.get("n_firms_dropped")),
            ("Treated × Post obs", did.get("n_did")),
        ]
        for k, v in did_rows:
            ws.cell(next_row, 1, k).font = Font(bold=True)
            ws.cell(next_row, 2, v)
            next_row += 1
        next_row += 1
    elif did_warning:
        ws.cell(next_row, 1, "DID skipped").font = Font(bold=True, color="C00000")
        ws.cell(next_row, 2, did_warning)
        next_row += 2

    start = next_row + 1
    _header(ws, start, ["Model", "Label", "Key Var", "p-value", "Sig",
                        "N obs", "R² Within", "R² Overall"])
    for i, row in enumerate(run_data["summary"], start + 1):
        ws.cell(i, 1, row["model"])
        ws.cell(i, 2, row["label"])
        ws.cell(i, 3, row["key_var"] or "—")
        ws.cell(i, 4, row["key_p"])
        ws.cell(i, 5, row["key_sig"])
        ws.cell(i, 6, row["n_obs"])
        ws.cell(i, 7, row["rsq_within"])
        ws.cell(i, 8, row["rsq_overall"])
        if row["key_sig"] in ("**", "***"):
            ws.cell(i, 5).font = Font(bold=True, color="C00000")
        elif row["key_sig"] == "*":
            ws.cell(i, 5).font = Font(bold=True, color="E26B0A")

    # Append errored models so they aren't silently missing
    errored = [(m, d) for m, d in run_data["results"].items() if "error" in d]
    if errored:
        err_start = start + 1 + len(run_data["summary"])
        for i, (m, d) in enumerate(errored, err_start):
            ws.cell(i, 1, m)
            ws.cell(i, 2, f"ERROR: {d['error']}").font = Font(color="C00000", italic=True)
            ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=8)

    for letter, w in zip("ABCDEFGH", [22, 28, 32, 12, 8, 10, 12, 12]):
        ws.column_dimensions[letter].width = w
    ws.freeze_panes = "A2"


def _write_model_sheet(ws, name, data):
    ws.merge_cells("A1:H1")
    c = ws.cell(
        1, 1,
        f"{name} — N={data['n_obs']:,}, R²w={data['rsq_within']:.4f}, "
        f"R²o={data['rsq_overall']:.4f}"
    )
    c.font = Font(bold=True, color="FFFFFF", size=12)
    c.fill = PatternFill("solid", fgColor=DARK_BLUE)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    _header(ws, 3, ["Variable", "Coef", "Std Error", "t-stat",
                    "p-value", "CI 95% lower", "CI 95% upper", "Sig"], bg=TEAL)
    for i, c in enumerate(data["coefs"], 4):
        bg = LIGHT_GRAY if i % 2 == 0 else None
        cells = [c["var"], round(c["coef"], 4), round(c["se"], 4),
                 round(c["tstat"], 3), round(c["pvalue"], 4),
                 round(c["ci_low"], 4), round(c["ci_high"], 4), c["sig"]]
        for ci, val in enumerate(cells, 1):
            cell = ws.cell(i, ci, val)
            if bg:
                cell.fill = PatternFill("solid", fgColor=bg)
            cell.alignment = Alignment(horizontal="left" if ci == 1 else "center")
            sig = c["sig"]
            if ci in (2, 5, 8):
                if sig in ("**", "***"):
                    cell.font = Font(bold=True, color="C00000")
                elif sig == "*":
                    cell.font = Font(bold=True, color="E26B0A")

    for letter, w in zip("ABCDEFGH", [32, 12, 12, 10, 12, 14, 14, 8]):
        ws.column_dimensions[letter].width = w
    ws.freeze_panes = "A4"


def _write_comparison(ws, run_data):
    all_vars = []
    for m, d in run_data["results"].items():
        if "error" in d:
            continue
        for c in d["coefs"]:
            if c["var"] not in all_vars:
                all_vars.append(c["var"])

    models = [m for m, d in run_data["results"].items() if "error" not in d]

    headers = ["Variable"]
    for m in models:
        headers += [f"{m} Coef", f"{m} p", "Sig"]
    _header(ws, 1, headers, bg=DARK_BLUE)

    for i, v in enumerate(all_vars, 2):
        ws.cell(i, 1, v).font = Font(bold=True)
        ws.cell(i, 1).alignment = Alignment(horizontal="left")
        for j, m in enumerate(models):
            d = run_data["results"][m]
            coef_data = next((c for c in d["coefs"] if c["var"] == v), None)
            base_col = 2 + j * 3
            if coef_data:
                p = coef_data["pvalue"]
                ws.cell(i, base_col, round(coef_data["coef"], 4))
                ws.cell(i, base_col + 1, round(p, 4))
                ws.cell(i, base_col + 2, coef_data["sig"])
                if p < 0.05:
                    for off in range(3):
                        ws.cell(i, base_col + off).font = Font(bold=True, color="C00000")
                elif p < 0.10:
                    for off in range(3):
                        ws.cell(i, base_col + off).font = Font(bold=True, color="E26B0A")
            else:
                for off in range(3):
                    ws.cell(i, base_col + off, "—")

    ws.column_dimensions["A"].width = 32
    for j in range(len(models)):
        for off, w in enumerate([12, 12, 8]):
            ws.column_dimensions[get_column_letter(2 + j * 3 + off)].width = w
    ws.freeze_panes = "B2"
